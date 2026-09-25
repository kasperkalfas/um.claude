"""Dane giełdowe do zadań 1–2: The Coca-Cola Company (NYSE: KO), kwartalnie.

Liczby są prawdziwe: raporty 10-Q i 10-K złożone w SEC (EDGAR, CIK 0000021344),
w mln USD. IV kwartał = rok (10-K) minus trzy kwartały narastająco (10-Q),
bo spółka nie raportuje IV kwartału osobno. Q2 2026: 10-Q z 29.07.2026.

Uruchomienie:  python generuj_gielda.py
Tworzy dwa pliki .xlsx obok skryptu i wypisuje klucz odpowiedzi do zadań 1 i 2.
Opcja --pobierz odświeża liczby prosto z SEC (wymaga internetu).
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TU = Path(__file__).parent

# kwartał: przychody, koszt sprzedaży, zysk brutto, koszty sprzedaży i ogólne (SG&A),
#          zysk operacyjny, zysk netto  [mln USD]
DANE = {
    "2021Q1": (9020, 3505, 5515, 2669, 2722, 2245),
    "2021Q2": (10129, 3787, 6342, 3017, 3016, 2641),
    "2021Q3": (10042, 3977, 6065, 3122, 2898, 2471),
    "2021Q4": (9464, 4088, 5376, 3336, 1672, 2414),
    "2022Q1": (10491, 4091, 6400, 2967, 3405, 2781),
    "2022Q2": (11325, 4830, 6495, 3203, 2341, 1905),
    "2022Q3": (11063, 4566, 6497, 3279, 3088, 2825),
    "2022Q4": (10125, 4513, 5612, 3431, 2075, 2031),
    "2023Q1": (10980, 4317, 6663, 3185, 3367, 3107),
    "2023Q2": (11972, 4912, 7060, 3321, 2401, 2547),
    "2023Q3": (11953, 4657, 7296, 3667, 3270, 3087),
    "2023Q4": (10849, 4634, 6215, 3799, 2273, 1973),
    "2024Q1": (11300, 4235, 7065, 3351, 2141, 3177),
    "2024Q2": (12363, 4812, 7551, 3549, 2632, 2411),
    "2024Q3": (11854, 4664, 7190, 3636, 2510, 2848),
    "2024Q4": (11544, 4613, 6931, 4046, 2709, 2195),
    "2025Q1": (11129, 4163, 6966, 3234, 3659, 3330),
    "2025Q2": (12535, 4714, 7821, 3470, 4280, 3810),
    "2025Q3": (12455, 4797, 7658, 3618, 3982, 3696),
    "2025Q4": (11822, 4723, 7099, 4199, 1841, 2271),
    "2026Q1": (12472, 4620, 7852, 3472, 4359, 3924),
    "2026Q2": (13380, 4965, 8415, 3720, 4672, 4425),
}

TAGI = ["Revenues", "CostOfGoodsAndServicesSold", "GrossProfit",
        "SellingGeneralAndAdministrativeExpense", "OperatingIncomeLoss", "NetIncomeLoss"]

NAGL = ["Kwartał", "Przychody", "Koszt sprzedaży", "Zysk brutto",
        "Koszty sprzedaży i ogólne (SG&A)", "Pozostałe koszty operacyjne",
        "Zysk operacyjny", "Zysk netto"]

BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
ZAMR = PatternFill("solid", fgColor="E7E6E6")
MLN = '#,##0'
PROC = '0.0%'


def pobierz():
    """Odświeża DANE z SEC (frames API, kwartały 3-miesięczne)."""
    import json
    import urllib.request
    ua = {"User-Agent": "Szkolenie Claude kontakt@example.com"}
    for k in DANE:
        rok, kw = k.split("Q")
        wiersz = []
        for tag in TAGI:
            url = f"https://data.sec.gov/api/xbrl/frames/us-gaap/{tag}/USD/CY{rok}Q{kw}.json"
            with urllib.request.urlopen(urllib.request.Request(url, headers=ua)) as r:
                dane = json.load(r)["data"]
            v = [x["val"] for x in dane if x["cik"] == 21344]
            wiersz.append(round(v[0] / 1e6) if v else None)
        if None not in wiersz:
            DANE[k] = tuple(wiersz)
        print(k, wiersz)


def etykieta(k):
    rok, kw = k.split("Q")
    return f"{['I', 'II', 'III', 'IV'][int(kw) - 1]} kw. {rok}"


def nastepny(k):
    rok, kw = int(k[:4]), int(k[-1])
    return f"{rok + (kw == 4)}Q{kw % 4 + 1}"


def poprzedni_rok(k):
    return f"{int(k[:4]) - 1}Q{k[-1]}"


def szerokosci(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def ttm(do, idx):
    """Suma 4 kwartałów kończących się na `do` (trailing twelve months)."""
    k, s = do, 0
    for _ in range(4):
        s += DANE[k][idx]
        k = f"{int(k[:4]) - (k[-1] == '1')}Q{(int(k[-1]) - 2) % 4 + 1}"
    return s


def zalozenia(do):
    """Założenia prognozy znane po zamknięciu kwartału `do`."""
    rok_wczesniej = f"{int(do[:4]) - 1}Q{do[-1]}"
    g = ttm(do, 0) / ttm(rok_wczesniej, 0) - 1
    mb = ttm(do, 2) / ttm(do, 0)
    sga = ttm(do, 3) / ttm(do, 0)
    return g, mb, sga


def prognoza(do, ile):
    """Prognoza `ile` kwartałów po `do`: ten sam kwartał rok wcześniej × (1+g)."""
    g, mb, sga = zalozenia(do)
    znane = {k: v[0] for k, v in DANE.items() if k <= do}
    wynik, k = [], do
    for _ in range(ile):
        k = nastepny(k)
        p = znane[poprzedni_rok(k)] * (1 + g)
        znane[k] = p
        zb = p * mb
        koszty = p * sga
        wynik.append((k, p, zb, koszty, zb - koszty))
    return wynik


def plik_dane():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dane"
    ws.append(NAGL)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (k, (p, ks, zb, sga, zo, zn)) in enumerate(DANE.items(), 2):
        ws.append([etykieta(k), p, ks, zb, sga, None, zo, zn])
        ws.cell(i, 6, f"=D{i}-E{i}-G{i}")
        for c in range(2, 9):
            ws.cell(i, c).number_format = MLN
    ws.freeze_panes = "B2"
    szerokosci(ws, [14, 12, 14, 12, 18, 16, 14, 12])

    op = wb.create_sheet("Opis")
    for w in [
        ["The Coca-Cola Company (NYSE: KO) – rachunek wyników kwartalnie, mln USD"],
        [],
        ["Skąd", "raporty kwartalne 10-Q i roczne 10-K złożone w SEC (EDGAR, CIK 0000021344) – dane publiczne"],
        ["IV kwartał", "rok z 10-K minus I–III kwartał narastająco (spółka nie raportuje IV kw. osobno)"],
        ["Ostatni kwartał", "II kw. 2026 (okres do 3.07.2026, raport z 29.07.2026)"],
        ["Pozostałe koszty operacyjne", "formuła: zysk brutto − SG&A − zysk operacyjny; odpisy, restrukturyzacje, zdarzenia jednorazowe"],
        ["Kwartały spółki", "kończą się w piątek najbliższy końcowi kwartału kalendarzowego"],
        [],
        ["To dane publiczne spółki giełdowej, nie dane Urzędu – wolno je wklejać do Claude."],
    ]:
        op.append(w)
    op["A1"].font = Font(bold=True, size=12)
    szerokosci(op, [26, 100])
    wb.save(TU / "coca-cola_kwartalnie_2021-2026.xlsx")


def plik_poprzednia_prognoza():
    g, mb, sga = zalozenia("2026Q1")
    wb = Workbook()
    ws = wb.active
    ws.title = "Prognoza"
    ws.append(["Prognoza krocząca sporządzona po I kw. 2026 (mln USD) – 6 kwartałów = 18 miesięcy"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Kwartał", "Status", "Przychody", "Zysk brutto", "SG&A", "Zysk operacyjny"])
    for c in ws[3]:
        c.font, c.fill = BOLD, HEAD
    for k in ["2025Q3", "2025Q4", "2026Q1"]:
        p, _, zb, s, zo, _ = DANE[k]
        ws.append([etykieta(k), "wykonanie (zamrożone)", p, zb, s, zo])
        for c in ws[ws.max_row]:
            c.fill = ZAMR
    for k, p, zb, s, zo in prognoza("2026Q1", 6):
        ws.append([etykieta(k), "prognoza", round(p), round(zb), round(s), round(zo)])
    for row in ws.iter_rows(min_row=4, min_col=3, max_col=6):
        for c in row:
            c.number_format = MLN
    szerokosci(ws, [14, 22, 12, 12, 12, 16])

    z = wb.create_sheet("Zalozenia")
    z.append(["Założenie", "Wartość", "Jak policzone (dane do I kw. 2026 włącznie)"])
    for c in z[1]:
        c.font, c.fill = BOLD, HEAD
    z.append(["Wzrost przychodów r/r", g, "przychody 4 ostatnich kwartałów ÷ 4 kwartały rok wcześniej − 1"])
    z.append(["Marża brutto", mb, "zysk brutto ÷ przychody, 4 ostatnie kwartały"])
    z.append(["SG&A jako % przychodów", sga, "SG&A ÷ przychody, 4 ostatnie kwartały"])
    z.append(["Pozostałe koszty operacyjne", 0, "założenie: brak zdarzeń jednorazowych"])
    z.append(["Metoda", None, "kwartał prognozy = ten sam kwartał rok wcześniej × (1 + wzrost r/r) – zachowuje sezonowość"])
    for r in range(2, 5):
        z.cell(r, 2).number_format = PROC
    szerokosci(z, [28, 10, 90])
    wb.save(TU / "prognoza_kroczaca_po_I_kw_2026.xlsx")


def klucz():
    rok = lambda r: sum(DANE[f"{r}Q{k}"][0] for k in range(1, 5))
    print("\n=== KLUCZ – ZADANIE 1 (financial-analyst, forecast_builder.py) ===")
    print(f"Przychody roczne: 2023 {rok(2023):,} | 2024 {rok(2024):,} | 2025 {rok(2025):,} mln USD")
    print(f"Ostatnie 4 kwartały (III kw. 2025 – II kw. 2026): {ttm('2026Q2', 0):,} mln USD")
    g, mb, sga = zalozenia("2026Q2")
    print(f"Założenia po II kw. 2026: wzrost r/r {g:.1%}, marża brutto {mb:.1%}, SG&A {sga:.1%}")
    ost = DANE["2026Q2"][0]
    mies = (1 + g) ** (1 / 12) - 1
    x, suma = ost, 0
    for _ in range(12):
        x *= 1 + mies
        suma += x
    print(f"Skrypt z domyślnymi 12 okresami: start od {ost:,} (II kw. = szczyt), "
          f"suma ≈ {suma:,.0f} mln USD – ok. {suma / ttm('2026Q2', 0):.1f}× za dużo "
          "(traktuje kwartały jak miesiące)")
    x, s4 = ost, 0
    for _ in range(4):
        x *= 1 + mies
        s4 += x
    sez = sum(p for _, p, *_ in prognoza("2026Q2", 4))
    print(f"Skrypt z forecast_periods=4: ≈ {s4:,.0f} mln USD (bez sezonowości, każdy kwartał jak II kw.)")
    print(f"Metoda sezonowa (kwartał rok wcześniej × (1+g)), III kw. 2026 – II kw. 2027: {sez:,.0f} mln USD")
    print("Sezonowość (średni udział kwartału w roku 2021–2025):",
          ", ".join(f"{['I', 'II', 'III', 'IV'][q - 1]} {sum(DANE[f'{r}Q{q}'][0] / rok(r) for r in range(2021, 2026)) / 5:.1%}"
                    for q in range(1, 5)))
    print(f"Zysk operacyjny, ostatnie 4 kwartały (z jednorazowymi): {ttm('2026Q2', 4):,} mln USD")
    pk = [(k, DANE[k][2] - DANE[k][3] - DANE[k][4]) for k in DANE]
    print("Największe pozostałe koszty operacyjne (jednorazowe):",
          ", ".join(f"{etykieta(k)} {v:,}" for k, v in sorted(pk, key=lambda t: -t[1])[:3]))

    print("\n=== KLUCZ – ZADANIE 2 (codexkit-fpa-rolling-forecast) ===")
    stara = {k: p for k, p, *_ in prognoza("2026Q1", 6)}
    stara_zo = {k: zo for k, *_, zo in prognoza("2026Q1", 6)}
    nowa = {k: p for k, p, *_ in prognoza("2026Q2", 6)}
    nowa_zo = {k: zo for k, *_, zo in prognoza("2026Q2", 6)}
    g1, mb1, sga1 = zalozenia("2026Q1")
    print(f"Założenia po I kw.: wzrost {g1:.1%}, marża {mb1:.1%}, SG&A {sga1:.1%}  →  "
          f"po II kw.: wzrost {g:.1%}, marża {mb:.1%}, SG&A {sga:.1%}")
    wyk = DANE["2026Q2"]
    print(f"II kw. 2026: prognoza {stara['2026Q2']:,.0f}, wykonanie {wyk[0]:,} → "
          f"{wyk[0] - stara['2026Q2']:+,.0f} ({wyk[0] / stara['2026Q2'] - 1:+.1%})")
    print(f"  zysk operacyjny: prognoza {stara_zo['2026Q2']:,.0f}, wykonanie {wyk[4]:,} "
          f"({wyk[4] - stara_zo['2026Q2']:+,.0f})")
    rok26_stary = DANE["2026Q1"][0] + stara["2026Q2"] + stara["2026Q3"] + stara["2026Q4"]
    rok26_nowy = DANE["2026Q1"][0] + wyk[0] + nowa["2026Q3"] + nowa["2026Q4"]
    print(f"Rok 2026: poprzednio {rok26_stary:,.0f}, teraz {rok26_nowy:,.0f} "
          f"({rok26_nowy - rok26_stary:+,.0f}, {rok26_nowy / rok26_stary - 1:+.1%})")
    d_q2 = wyk[0] - stara["2026Q2"]
    d_h2 = nowa["2026Q3"] + nowa["2026Q4"] - stara["2026Q3"] - stara["2026Q4"]
    print(f"  most: II kw. wykonanie ponad prognozę {d_q2:+,.0f} + wyższy wzrost w III–IV kw. {d_h2:+,.0f}")
    print("Nowe okno III kw. 2026 – IV kw. 2027 (przychody / zysk operacyjny):")
    for k in nowa:
        print(f"  {etykieta(k):13} {nowa[k]:>8,.0f} {nowa_zo[k]:>8,.0f}")
    print("Kubełki Volume / Price-Mix / FX: BRAK w danych – Claude ma to napisać, nie zgadywać.")


if __name__ == "__main__":
    if "--pobierz" in sys.argv:
        pobierz()
    plik_dane()
    plik_poprzednia_prognoza()
    klucz()
