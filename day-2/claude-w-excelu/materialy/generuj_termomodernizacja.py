"""Tworzy plik wejściowy do zadania 9 (model finansowy inwestycji) i wypisuje
klucz odpowiedzi. Dane w 100% fikcyjne.

Uruchomienie:  python generuj_termomodernizacja.py
Wynik:         termomodernizacja_zalozenia.xlsx (jeden arkusz `Zalozenia`)
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding="utf-8")  # polskie znaki w konsoli Windows

# --- założenia (wszystkie fikcyjne) ---------------------------------------
NAKLAD = 4_200_000        # nakład inwestycyjny brutto, zł
DOFINANSOWANIE = 0.45     # udział dotacji zewnętrznej w nakładzie
OSZCZEDNOSC_R1 = 310_000  # oszczędność kosztów energii w roku 1, zł
WZROST_CEN = 0.04         # roczny wzrost cen energii
UTRZYMANIE_R1 = 15_000    # koszt serwisu nowych instalacji w roku 1, zł
WZROST_UTRZ = 0.03        # roczny wzrost kosztów serwisu
HORYZONT = 15             # lat
STOPA = 0.06              # stopa dyskontowa (koszt długu gminy)
REZYDUALNA = 0.20         # wartość rezydualna po horyzoncie, % nakładu
KREDYT_OPROC = 0.055      # oprocentowanie kredytu na wkład własny
KREDYT_LATA = 10          # okres kredytu, raty roczne równe (annuitetowe)

wb = Workbook()
ws = wb.active
ws.title = "Zalozenia"

bold = Font(bold=True)
head = PatternFill("solid", fgColor="D9E1F2")

ws["A1"] = "Termomodernizacja budynku Szkoły Podstawowej nr 99 (dane fikcyjne)"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = "Założenia do oceny opłacalności i finansowania – Wydział Finansowy"

rows = [
    ("Pozycja", "Wartość", "Jednostka", "Uwagi"),
    ("Nakład inwestycyjny brutto", NAKLAD, "zł", "kosztorys inwestorski"),
    ("Dofinansowanie zewnętrzne", DOFINANSOWANIE, "% nakładu", "umowa o dofinansowanie"),
    ("Oszczędność kosztów energii w roku 1", OSZCZEDNOSC_R1, "zł/rok", "audyt energetyczny"),
    ("Roczny wzrost cen energii", WZROST_CEN, "%", "założenie"),
    ("Koszt serwisu nowych instalacji w roku 1", UTRZYMANIE_R1, "zł/rok", "oferta wykonawcy"),
    ("Roczny wzrost kosztów serwisu", WZROST_UTRZ, "%", "założenie"),
    ("Horyzont analizy", HORYZONT, "lat", ""),
    ("Stopa dyskontowa", STOPA, "%", "koszt długu gminy"),
    ("Wartość rezydualna po horyzoncie", REZYDUALNA, "% nakładu", "założenie"),
    ("Kredyt – oprocentowanie", KREDYT_OPROC, "%", "oferta banku"),
    ("Kredyt – okres spłaty", KREDYT_LATA, "lat", "raty roczne równe"),
]
start = 4
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = ws.cell(row=start + i, column=j + 1, value=v)
        if i == 0:
            c.font = bold
            c.fill = head
        elif j == 1 and isinstance(v, float) and v < 1:
            c.number_format = "0.0%"
        elif j == 1 and isinstance(v, int) and v > 100:
            c.number_format = "#,##0"
ws["A18"] = ("Uwaga: wkład własny gminy (nakład minus dofinansowanie) ma być "
             "sfinansowany kredytem. Wszystkie kwoty i podmioty są fikcyjne.")
ws["A18"].alignment = Alignment(wrap_text=True)
for col, w in zip("ABCD", (42, 14, 12, 28)):
    ws.column_dimensions[col].width = w

wb.save("termomodernizacja_zalozenia.xlsx")
print("Zapisano termomodernizacja_zalozenia.xlsx")

# --- klucz odpowiedzi -------------------------------------------------------
def model(r=STOPA, g=WZROST_CEN, dof=DOFINANSOWANIE):
    wklad = NAKLAD * (1 - dof)
    pv = 0.0
    cum = -wklad
    payback = None
    for t in range(1, HORYZONT + 1):
        net = OSZCZEDNOSC_R1 * (1 + g) ** (t - 1) - UTRZYMANIE_R1 * (1 + WZROST_UTRZ) ** (t - 1)
        pv += net / (1 + r) ** t
        cum += net
        if payback is None and cum >= 0:
            payback = t
    pv_rez = NAKLAD * REZYDUALNA / (1 + r) ** HORYZONT
    return wklad, pv, pv_rez, -wklad + pv + pv_rez, payback

wklad, pv, pv_rez, npv, pb = model()
print(f"\nKLUCZ (stopa {STOPA:.0%}, wzrost cen {WZROST_CEN:.0%}, dotacja {DOFINANSOWANIE:.0%})")
print(f"  wkład własny (kredyt):        {wklad:>12,.0f} zł")
print(f"  PV oszczędności netto 1–15:   {pv:>12,.0f} zł")
print(f"  PV wartości rezydualnej:      {pv_rez:>12,.0f} zł")
print(f"  NPV dla gminy:                {npv:>12,.0f} zł")
print(f"  prosty okres zwrotu wkładu:   {pb} lat")
lo, hi = 0.0, 0.5
for _ in range(60):
    mid = (lo + hi) / 2
    lo, hi = (mid, hi) if model(r=mid)[3] > 0 else (lo, mid)
print(f"  IRR:                          {lo:.1%}")
print(f"  NPV bez dofinansowania:       {model(dof=0)[3]:>12,.0f} zł")
print("  wrażliwość NPV (stopa × wzrost cen):")
for r in (0.04, 0.06, 0.08, 0.10):
    print("   ", f"{r:.0%}", "  ".join(f"{model(r=r, g=g)[3]:>11,.0f}" for g in (0.0, 0.02, 0.04, 0.06)))
A = wklad * KREDYT_OPROC / (1 - (1 + KREDYT_OPROC) ** -KREDYT_LATA)
print(f"  rata roczna kredytu (PMT):    {A:>12,.0f} zł, odsetki razem {A * KREDYT_LATA - wklad:,.0f} zł")
