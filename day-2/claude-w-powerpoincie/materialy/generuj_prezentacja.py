"""Tworzy celowo przeciętną prezentację `wykonanie_budzetu_2026_8m.pptx`
(materiał do zadania 1 – poprawianie gotowej prezentacji z Claude) i
wypisuje klucz: wszystkie liczby, które muszą przetrwać przeróbkę.

Dane: `../../../day-1/claude-code-cli/materialy/zestawienie_roczne_2026.xlsx`
(fikcyjne wykonanie wydatków I–VIII 2026 wg działów). Prognoza roczna =
średnia miesięczna × 12.

Uruchomienie:  python generuj_prezentacja.py
"""
import sys
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.util import Pt, Inches

sys.stdout.reconfigure(encoding="utf-8")

SRC = Path(__file__).resolve().parents[3] / "day-1" / "claude-code-cli" / "materialy" / "zestawienie_roczne_2026.xlsx"
OUT = Path(__file__).with_name("wykonanie_budzetu_2026_8m.pptx")
MIESIACE = 8

# --- dane -------------------------------------------------------------------
ws = openpyxl.load_workbook(SRC, data_only=True)["Wykonanie"]
dzialy = []
for r in ws.iter_rows(min_row=6, max_row=13, values_only=True):
    plan = r[2]
    wyk = sum(r[3:3 + MIESIACE])
    prog = wyk / MIESIACE * 12
    dzialy.append(dict(kod=r[0], nazwa=r[1], plan=plan, wyk=wyk, proc=wyk / plan,
                       prog=prog, odch=prog - plan, odch_proc=(prog - plan) / plan))
plan_r = sum(d["plan"] for d in dzialy)
wyk_r = sum(d["wyk"] for d in dzialy)
prog_r = wyk_r / MIESIACE * 12


def zl(x):
    return f"{x:,.0f}".replace(",", " ") + " zł"


def pc(x):
    return f"{x * 100:.1f} %".replace(".", ",")


def pcs(x):
    return f"{x * 100:+.1f} %".replace(".", ",")


# --- prezentacja (celowo: domyślny szablon, dużo tekstu, zero grafiki) --------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

# 1. tytuł
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Informacja o wykonaniu budżetu wydatków za 8 miesięcy 2026 r."
s.placeholders[1].text = "Wydział Finansowy – Urząd Miejski w Opolu\nmateriał na posiedzenie Komisji Budżetowej, wrzesień 2026 (dane fikcyjne)"

# 2. ściana tekstu
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Wykonanie wydatków według działów"
tf = s.placeholders[1].text_frame
tf.word_wrap = True
zdania = []
for d in dzialy:
    zdania.append(
        f"W dziale {d['kod']} {d['nazwa']} plan roczny wynosi {zl(d['plan'])}, "
        f"wykonanie po ośmiu miesiącach {zl(d['wyk'])}, to jest {pc(d['proc'])} planu; "
        f"prognoza wykonania rocznego przy utrzymaniu średniego miesięcznego tempa wydatków "
        f"wynosi {zl(d['prog'])}, co oznacza odchylenie od planu {pcs(d['odch_proc'])}."
    )
zdania.append(
    f"Łącznie plan wydatków wynosi {zl(plan_r)}, wykonanie {zl(wyk_r)} ({pc(wyk_r / plan_r)} "
    f"przy upływie {pc(MIESIACE / 12)} roku), prognoza roczna {zl(prog_r)}, "
    f"czyli {pcs((prog_r - plan_r) / plan_r)} względem planu."
)
tf.text = " ".join(zdania)
for p in tf.paragraphs:
    for run in p.runs:
        run.font.size = Pt(12)

# 3. tabela
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Tabela: plan, wykonanie, prognoza"
rows, cols = len(dzialy) + 2, 6
tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5)).table
for j, h in enumerate(["Dział", "Nazwa", "Plan roczny", "Wykonanie I–VIII", "% planu", "Prognoza roczna"]):
    tbl.cell(0, j).text = h
for i, d in enumerate(dzialy, start=1):
    for j, v in enumerate([d["kod"], d["nazwa"], zl(d["plan"]), zl(d["wyk"]), pc(d["proc"]), zl(d["prog"])]):
        tbl.cell(i, j).text = v
for j, v in enumerate(["", "Razem", zl(plan_r), zl(wyk_r), pc(wyk_r / plan_r), zl(prog_r)]):
    tbl.cell(rows - 1, j).text = v
for i in range(rows):
    for j in range(cols):
        for p in tbl.cell(i, j).text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(11)

# 4. prognoza – wypunktowanie
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Prognoza wykonania rocznego"
tf = s.placeholders[1].text_frame
tf.text = f"Prognoza roczna ogółem: {zl(prog_r)} wobec planu {zl(plan_r)} ({pcs((prog_r - plan_r) / plan_r)})"
for d in sorted(dzialy, key=lambda d: d["odch_proc"]):
    p = tf.add_paragraph()
    p.text = f"{d['kod']} {d['nazwa']}: {pcs(d['odch_proc'])} ({zl(d['odch'])})"
    p.level = 1

# 5. wnioski
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Wnioski i rekomendacje"
tf = s.placeholders[1].text_frame
tf.text = "Wykonanie wydatków przebiega zgodnie z planem."
for t in [
    "W dwóch działach prognoza przekracza plan – wymagana analiza.",
    "W pozostałych działach prognoza jest niższa od planu.",
    "Rekomenduje się monitorowanie wykonania w kolejnych miesiącach.",
    "Materiał zostanie zaktualizowany po zamknięciu września.",
]:
    tf.add_paragraph().text = t

prs.save(OUT)
print(f"Zapisano {OUT.name} (5 slajdów)")

# --- klucz ------------------------------------------------------------------
print("\nKLUCZ – liczby, które muszą przetrwać przeróbkę slajdu 2 (i całej prezentacji):")
print(f"{'dział':<6}{'plan':>15}{'wykonanie':>15}{'% planu':>9}{'prognoza':>15}{'odchylenie':>12}")
for d in dzialy:
    print(f"{d['kod']:<6}{d['plan']:>15,}{d['wyk']:>15,}{d['proc']*100:>8.1f}%{d['prog']:>15,.0f}{d['odch_proc']*100:>+11.1f}%")
print(f"{'RAZEM':<6}{plan_r:>15,}{wyk_r:>15,}{wyk_r/plan_r*100:>8.1f}%{prog_r:>15,.0f}{(prog_r-plan_r)/plan_r*100:>+11.1f}%")
print(f"\nUpływ roku: {MIESIACE}/12 = {MIESIACE/12*100:.1f} %. Wykonanie ogółem {wyk_r/plan_r*100:.1f} % – "
      f"{'poniżej' if wyk_r/plan_r < MIESIACE/12 else 'powyżej'} proporcji.")
nad = [d for d in dzialy if d["odch"] > 0]
print(f"Działy z prognozą POWYŻEJ planu: {', '.join(d['kod'] + ' (' + pcs(d['odch_proc']) + ')' for d in nad)}")
print(f"Największe niedowykonanie: {min(dzialy, key=lambda d: d['odch_proc'])['kod']} "
      f"({pcs(min(d['odch_proc'] for d in dzialy))}), największa kwota: "
      f"{min(dzialy, key=lambda d: d['odch'])['kod']} ({zl(min(d['odch'] for d in dzialy))})")
print(f"Slajd 2 zawiera {len(dzialy) * 5 + 6} liczb (po 5 na dział + 6 w zdaniu „Łącznie”), nie licząc kodów działów.")
print("\nPrzykładowy tytuł-teza (DataPOV): "
      f"„Budżet 2026 po 8 miesiącach: wykonanie {pcs((prog_r-plan_r)/plan_r)} do planu, "
      f"{len(nad)} działy wymagają korekty planu”")
