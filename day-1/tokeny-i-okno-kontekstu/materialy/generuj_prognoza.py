"""Generator fikcyjnych danych do zadania 9 (prognoza wykonania z zewnetrznym skillem).

Uruchom:  python generuj_prognoza.py
Tworzy (nadpisuje) w biezacym folderze:
  wykonanie_miesieczne_2023-2025.xlsx  - 36 miesiecy x 8 dzialow, dane "historyczne" (uczenie)
  kalendarz_2026.xlsx                  - 12 miesiecy 2026 x 8 dzialow, puste kolumny do prognozy
Dane sa calkowicie wymyslone. Sezonowosc i trend sa celowe (patrz KLUCZ na koncu).
"""
import math
import random
from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(2026)

DZIALY = [
    ("600", "Transport i łączność",                      1_900_000, 0.04, {1: 1.25, 2: 1.20, 12: 1.15, 7: 0.90, 8: 0.90}),
    ("750", "Administracja publiczna",                    1_500_000, 0.02, {12: 1.10}),
    ("801", "Oświata i wychowanie",                       5_000_000, 0.05, {7: 0.70, 8: 0.70, 9: 1.15, 1: 1.05}),
    ("851", "Ochrona zdrowia",                              270_000, 0.03, {1: 1.10, 2: 1.10, 11: 1.05}),
    ("852", "Pomoc społeczna",                            1_700_000, 0.06, {12: 1.20, 1: 1.05}),
    ("900", "Gospodarka komunalna i ochrona środowiska", 1_250_000, 0.03, {5: 1.10, 6: 1.10, 9: 1.05}),
    ("921", "Kultura i ochrona dziedzictwa narodowego",    600_000, 0.02, {6: 1.30, 7: 1.20, 8: 1.20, 12: 1.10}),
    ("926", "Kultura fizyczna",                            430_000, 0.03, {6: 1.25, 7: 1.25, 8: 1.15, 1: 0.85}),
]

def wykonanie(baza, trend_roczny, sezon, rok_idx, miesiac):
    poziom = baza * (1 + trend_roczny) ** rok_idx
    s = sezon.get(miesiac, 1.0)
    szum = random.gauss(1.0, 0.04)
    return int(round(poziom * s * szum, -3))

wb = Workbook()
ws = wb.active
ws.title = "Wykonanie"
ws.append(["Miesiąc", "Dział", "Nazwa działu", "Wykonanie"])
for c in ws[1]:
    c.font = Font(bold=True)
for rok_idx, rok in enumerate((2023, 2024, 2025)):
    for m in range(1, 13):
        for kod, nazwa, baza, trend, sezon in DZIALY:
            ws.append([f"{rok}-{m:02d}", kod, nazwa, wykonanie(baza, trend, sezon, rok_idx, m)])
ws.column_dimensions["A"].width = 10
ws.column_dimensions["C"].width = 44
ws.column_dimensions["D"].width = 14
wb.save("wykonanie_miesieczne_2023-2025.xlsx")

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Kalendarz_2026"
ws2.append(["Miesiąc", "Dział", "Nazwa działu", "Prognoza", "Dolna granica 95%", "Górna granica 95%"])
for c in ws2[1]:
    c.font = Font(bold=True)
for m in range(1, 13):
    for kod, nazwa, *_ in DZIALY:
        ws2.append([f"2026-{m:02d}", kod, nazwa, None, None, None])
for col, w in zip("ABCDEF", (10, 8, 44, 14, 18, 18)):
    ws2.column_dimensions[col].width = w
wb2.save("kalendarz_2026.xlsx")

print("Zapisano: wykonanie_miesieczne_2023-2025.xlsx, kalendarz_2026.xlsx")
print("\nKLUCZ (co powinien wykryc dobry model):")
for kod, nazwa, baza, trend, sezon in DZIALY:
    szczyty = ", ".join(f"{m}:{v:.2f}" for m, v in sorted(sezon.items()))
    print(f"  {kod} {nazwa[:28]:28} trend +{trend*100:.0f}%/rok  sezon (miesiac:mnoznik) {szczyty}")
print("\nSprawdzian dla uczestnikow: prosta prognoza 2026 = srednia 2025 x 12 x (1 + trend)")
