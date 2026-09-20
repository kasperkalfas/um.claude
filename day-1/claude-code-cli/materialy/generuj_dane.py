"""Generuje fikcyjne pliki wsadowe do zadan z Claude Code (Dzien 1, claude-code-cli).

Uruchom: python generuj_dane.py
Nadpisuje wszystkie pliki w tym folderze (poza tym skryptem) - uzyj, gdy
uczestnicy zmodyfikowali materialy i potrzebujesz czystej wersji.
"""
import csv
import os
import random
import shutil
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
# Konsola Windows bez UTF-8 (cp852/cp1250) wysypuje print() z polskimi znakami.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
random.seed(2026)

DZIALY = [
    ("600", "Transport i łączność", 24_000_000),
    ("750", "Administracja publiczna", 18_500_000),
    ("801", "Oświata i wychowanie", 62_000_000),
    ("851", "Ochrona zdrowia", 3_200_000),
    ("852", "Pomoc społeczna", 21_000_000),
    ("900", "Gospodarka komunalna i ochrona środowiska", 15_800_000),
    ("921", "Kultura i ochrona dziedzictwa narodowego", 7_400_000),
    ("926", "Kultura fizyczna", 5_600_000),
]
MIESIACE = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
PLN = '#,##0 "zł"'


def naglowek(ws, row, values):
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def szerokosci(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --- wykonanie I-VIII (deterministyczne) ------------------------------------
WYKONANIE = {}
for kod, _, plan in DZIALY:
    baza = plan / 12
    WYKONANIE[kod] = [round(baza * random.uniform(0.82, 1.12) / 1000) * 1000 for _ in range(8)]

# --- eksport ERP: wrzesien (brudny) i pazdziernik (czysty) ------------------
ERP = {
    "2026-09": {
        "600": [("60016", "4270", "Remont nawierzchni ul. Niemodlińskiej", "Drogbud Opole Sp. z o.o.", 1_240_000, "FV/2026/09/0142"),
                ("60016", "4300", "Utrzymanie sygnalizacji świetlnej", "Signal-Tech S.A.", 310_000, "FV/2026/09/0151"),
                ("60004", "4300", "Dopłata do transportu zbiorowego", "MZK Opole Sp. z o.o.", 600_000, "NK/2026/09/0031")],
        "750": [("75023", "4010", "Wynagrodzenia pracowników urzędu", "lista płac", 980_000, "LP/2026/09"),
                ("75023", "4110", "Składki na ubezpieczenia społeczne", "ZUS", 190_000, "DRA/2026/09"),
                ("75023", "4260", "Energia elektryczna – budynek Rynek", "Tauron Sprzedaż", 62_000, "FV/2026/09/0203"),
                ("75023", "4210", "Materiały biurowe", "Papirus Hurt", 18_000, "FV/2026/09/0210"),
                ("75075", "4300", "Promocja miasta – kampania jesienna", "Agencja Kreatywna Nysa", 230_000, "FV/2026/09/0218")],
        "801": [("80101", "4010", "Wynagrodzenia – szkoły podstawowe", "lista płac", 3_100_000, "LP/2026/09/SP"),
                ("80104", "4010", "Wynagrodzenia – przedszkola", "lista płac", 1_250_000, "LP/2026/09/PP"),
                ("80101", "4260", "Energia – szkoły podstawowe", "Tauron Sprzedaż", 410_000, "FV/2026/09/0221"),
                ("80148", "4220", "Zakup żywności – stołówki", "Catering Odra", 450_000, "FV/2026/09/0230")],
        "851": [("85154", "4300", "Programy profilaktyczne", "Fundacja Zdrowe Opole", 165_000, "FV/2026/09/0240"),
                ("85195", "4210", "Wyposażenie gabinetów", "MedSupply Polska", 100_000, "FV/2026/09/0244")],
        "852": [("85219", "4010", "Wynagrodzenia – MOPR", "lista płac", 1_050_000, "LP/2026/09/MOPR"),
                ("85214", "3110", "Zasiłki okresowe i celowe", "świadczenia", 520_000, "ZAS/2026/09"),
                ("85228", "4300", "Usługi opiekuńcze", "Spółdzielnia Socjalna Pomocna Dłoń", 150_000, "FV/2026/09/0250")],
        "900": [("90015", "4260", "Oświetlenie ulic", "Tauron Dystrybucja", 720_000, "FV/2026/09/0260"),
                ("90003", "4300", "Oczyszczanie miasta", "Remondis Opole", 480_000, "FV/2026/09/0263"),
                ("90004", "4300", "Utrzymanie zieleni", "Zieleń Miejska Sp. z o.o.", 140_000, "FV/2026/09/0270")],
        "921": [("92109", "2480", "Dotacja – Miejski Ośrodek Kultury", "MOK Opole", 380_000, "DOT/2026/09/01"),
                ("92116", "2480", "Dotacja – Miejska Biblioteka Publiczna", "MBP Opole", 230_000, "DOT/2026/09/02")],
        "926": [("92601", "4270", "Remont hali sportowej", "Budosport Sp. z o.o.", 330_000, "FV/2026/09/0281"),
                ("92605", "2820", "Dotacje dla klubów sportowych", "kluby sportowe", 125_000, "DOT/2026/09/03")],
    },
    "2026-10": {
        "600": [("60016", "4270", "Remont chodników – Zaodrze", "Drogbud Opole Sp. z o.o.", 1_070_000, "FV/2026/10/0301"),
                ("60016", "4300", "Utrzymanie sygnalizacji świetlnej", "Signal-Tech S.A.", 310_000, "FV/2026/10/0305"),
                ("60004", "4300", "Dopłata do transportu zbiorowego", "MZK Opole Sp. z o.o.", 600_000, "NK/2026/10/0032")],
        "750": [("75023", "4010", "Wynagrodzenia pracowników urzędu", "lista płac", 985_000, "LP/2026/10"),
                ("75023", "4110", "Składki na ubezpieczenia społeczne", "ZUS", 191_000, "DRA/2026/10"),
                ("75023", "4260", "Energia elektryczna – budynek Rynek", "Tauron Sprzedaż", 71_000, "FV/2026/10/0310"),
                ("75023", "4210", "Materiały biurowe", "Papirus Hurt", 13_000, "FV/2026/10/0312"),
                ("75075", "4300", "Promocja miasta – targi", "Agencja Kreatywna Nysa", 250_000, "FV/2026/10/0318")],
        "801": [("80101", "4010", "Wynagrodzenia – szkoły podstawowe", "lista płac", 3_120_000, "LP/2026/10/SP"),
                ("80104", "4010", "Wynagrodzenia – przedszkola", "lista płac", 1_260_000, "LP/2026/10/PP"),
                ("80101", "4260", "Energia – szkoły podstawowe", "Tauron Sprzedaż", 500_000, "FV/2026/10/0321"),
                ("80148", "4220", "Zakup żywności – stołówki", "Catering Odra", 460_000, "FV/2026/10/0330")],
        "851": [("85154", "4300", "Programy profilaktyczne", "Fundacja Zdrowe Opole", 160_000, "FV/2026/10/0340"),
                ("85195", "4210", "Wyposażenie gabinetów", "MedSupply Polska", 80_000, "FV/2026/10/0344")],
        "852": [("85219", "4010", "Wynagrodzenia – MOPR", "lista płac", 1_055_000, "LP/2026/10/MOPR"),
                ("85214", "3110", "Zasiłki okresowe i celowe", "świadczenia", 490_000, "ZAS/2026/10"),
                ("85228", "4300", "Usługi opiekuńcze", "Spółdzielnia Socjalna Pomocna Dłoń", 145_000, "FV/2026/10/0350")],
        "900": [("90015", "4260", "Oświetlenie ulic", "Tauron Dystrybucja", 790_000, "FV/2026/10/0360"),
                ("90003", "4300", "Oczyszczanie miasta", "Remondis Opole", 485_000, "FV/2026/10/0363"),
                ("90004", "4300", "Utrzymanie zieleni – nasadzenia jesienne", "Zieleń Miejska Sp. z o.o.", 140_000, "FV/2026/10/0370")],
        "921": [("92109", "2480", "Dotacja – Miejski Ośrodek Kultury", "MOK Opole", 370_000, "DOT/2026/10/01"),
                ("92116", "2480", "Dotacja – Miejska Biblioteka Publiczna", "MBP Opole", 215_000, "DOT/2026/10/02")],
        "926": [("92601", "4270", "Remont hali sportowej – etap II", "Budosport Sp. z o.o.", 355_000, "FV/2026/10/0381"),
                ("92605", "2820", "Dotacje dla klubów sportowych", "kluby sportowe", 125_000, "DOT/2026/10/03")],
    },
}


def kwota_brudna(i, kwota):
    """Rozne formaty kwot, jak z roznych eksportow ERP."""
    warianty = [
        f"{kwota:,.2f}".replace(",", " ").replace(".", ","),  # 1 240 000,00
        f"{kwota:.2f}",                                        # 1240000.00
        f"{kwota:,.2f}".replace(",", " ").replace(".", ",") + " zł",
        str(kwota),
    ]
    return warianty[i % len(warianty)]


def zapisz_erp(miesiac, brudny):
    sciezka = os.path.join(HERE, f"eksport_erp_{miesiac}.csv")
    rok, mies = miesiac.split("-")
    with open(sciezka, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Data księgowania", "Dział", "Rozdział", "Paragraf", "Nazwa", "Kontrahent", "Kwota", "Nr dokumentu"])
        i = 0
        for kod, pozycje in ERP[miesiac].items():
            for rozdz, par, nazwa, kontr, kwota, nr in pozycje:
                dzien = 3 + (i * 7) % 25
                if brudny:
                    data = f"{dzien:02d}.{mies}.{rok}" if i % 3 == 0 else f"{rok}-{mies}-{dzien:02d}"
                    kw = kwota_brudna(i, kwota)
                    nazwa_out = nazwa + ("  " if i % 5 == 0 else "")
                    par_out = "" if nr == "FV/2026/09/0210" else par
                else:
                    data, kw, nazwa_out, par_out = f"{rok}-{mies}-{dzien:02d}", f"{kwota:.2f}", nazwa, par
                w.writerow([data, kod, rozdz, par_out, nazwa_out, kontr, kw, nr])
                if brudny and nr == "FV/2026/09/0142":
                    w.writerow([data, kod, rozdz, par_out, nazwa_out, kontr, kw, nr])  # duplikat
                i += 1


# --- zestawienie miesieczne: szablon ----------------------------------------
def zestawienie_miesieczne_szablon():
    wb = Workbook()
    ws = wb.active
    ws.title = "Zestawienie"
    ws["A1"] = "Zestawienie miesięczne wykonania wydatków – Urząd Miejski w Opolu (dane fikcyjne)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"], ws["A3"] = "Miesiąc:", "Rok:"
    ws["B3"] = 2026
    naglowek(ws, 5, ["Dział", "Nazwa działu", "Plan roczny", "Wykonanie w miesiącu",
                     "Wykonanie narastająco", "% wykonania planu"])
    for r, (kod, nazwa, plan) in enumerate(DZIALY, start=6):
        ws.cell(row=r, column=1, value=kod)
        ws.cell(row=r, column=2, value=nazwa)
        ws.cell(row=r, column=3, value=plan).number_format = PLN
        ws.cell(row=r, column=4).number_format = PLN
        ws.cell(row=r, column=5).number_format = PLN
        ws.cell(row=r, column=6, value=f'=IF(C{r}=0,"",E{r}/C{r})').number_format = "0.0%"
    ws["A14"], ws["A14"].font = "Razem", BOLD
    for col in "CDE":
        ws[f"{col}14"] = f"=SUM({col}6:{col}13)"
        ws[f"{col}14"].number_format = PLN
        ws[f"{col}14"].font = BOLD
    ws["F14"] = '=IF(C14=0,"",E14/C14)'
    ws["F14"].number_format = "0.0%"
    szerokosci(ws, [8, 42, 16, 20, 20, 16])
    wb.save(os.path.join(HERE, "zestawienie_miesieczne_SZABLON.xlsx"))


# --- zestawienie roczne -----------------------------------------------------
def zestawienie_roczne():
    wb = Workbook()
    ws = wb.active
    ws.title = "Wykonanie"
    ws["A1"] = "Wykonanie wydatków 2026 wg działów – Urząd Miejski w Opolu (dane fikcyjne)"
    ws["A1"].font = Font(bold=True, size=12)
    naglowek(ws, 5, ["Dział", "Nazwa działu", "Plan roczny"] + MIESIACE + ["Razem", "% planu"])
    for r, (kod, nazwa, plan) in enumerate(DZIALY, start=6):
        ws.cell(row=r, column=1, value=kod)
        ws.cell(row=r, column=2, value=nazwa)
        ws.cell(row=r, column=3, value=plan).number_format = PLN
        for m in range(12):
            c = ws.cell(row=r, column=4 + m)
            if m < 8:
                c.value = WYKONANIE[kod][m]
            c.number_format = PLN
        ws.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})").number_format = PLN
        ws.cell(row=r, column=17, value=f'=IF(C{r}=0,"",P{r}/C{r})').number_format = "0.0%"
    ws["A14"], ws["A14"].font = "Razem", BOLD
    for col_idx in range(3, 17):
        col = get_column_letter(col_idx)
        ws[f"{col}14"] = f"=SUM({col}6:{col}13)"
        ws[f"{col}14"].number_format = PLN
        ws[f"{col}14"].font = BOLD
    ws["Q14"] = '=IF(C14=0,"",P14/C14)'
    ws["Q14"].number_format = "0.0%"
    szerokosci(ws, [8, 42, 15] + [13] * 12 + [15, 10])
    ws.freeze_panes = "D6"

    p = wb.create_sheet("Prognoza")
    p["A1"] = "Prognoza wykonania rocznego – materiał dla banku (dane fikcyjne)"
    p["A1"].font = Font(bold=True, size=12)
    p["A2"] = "Liczba miesięcy z danymi:"
    p["B2"] = 8
    naglowek(p, 5, ["Dział", "Nazwa działu", "Plan roczny", "Wykonanie narastająco",
                    "Średnia miesięczna", "Prognoza roczna", "Odchylenie od planu"])
    for r, (kod, nazwa, _) in enumerate(DZIALY, start=6):
        p.cell(row=r, column=1, value=kod)
        p.cell(row=r, column=2, value=nazwa)
        p.cell(row=r, column=3, value=f"=Wykonanie!C{r}").number_format = PLN
        p.cell(row=r, column=4, value=f"=Wykonanie!P{r}").number_format = PLN
        p.cell(row=r, column=5, value=f"=D{r}/$B$2").number_format = PLN
        p.cell(row=r, column=6, value=f"=E{r}*12").number_format = PLN
        p.cell(row=r, column=7, value=f"=F{r}-C{r}").number_format = PLN
    p["A14"], p["A14"].font = "Razem", BOLD
    for col in "CDEFG":
        p[f"{col}14"] = f"=SUM({col}6:{col}13)"
        p[f"{col}14"].number_format = PLN
        p[f"{col}14"].font = BOLD
    szerokosci(p, [8, 42, 15, 20, 18, 16, 18])
    wb.save(os.path.join(HERE, "zestawienie_roczne_2026.xlsx"))


# --- zestawienie z bledami (zadanie 4) --------------------------------------
def zestawienie_bledy():
    wb = Workbook()
    ws = wb.active
    ws.title = "Zestawienie"
    ws["A1"] = "Zestawienie miesięczne wykonania wydatków – sierpień 2026 (dane fikcyjne)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"], ws["B2"] = "Miesiąc:", "sierpień 2026"
    ws["A3"], ws["B3"] = "Rok:", 2026
    naglowek(ws, 5, ["Dział", "Nazwa działu", "Plan roczny", "Wykonanie w miesiącu",
                     "Wykonanie narastająco", ""])  # BLAD: brak naglowka F5
    wiersze = []
    for kod, nazwa, plan in DZIALY:
        wiersze.append((kod, nazwa, plan, WYKONANIE[kod][7], sum(WYKONANIE[kod])))
    wiersze.insert(3, None)  # BLAD: pusty wiersz 9
    for r, w in enumerate(wiersze, start=6):
        if w is None:
            continue
        kod, nazwa, plan, mies, naras = w
        ws.cell(row=r, column=1, value=kod)
        ws.cell(row=r, column=2, value=nazwa)
        ws.cell(row=r, column=3, value=plan).number_format = PLN
        ws.cell(row=r, column=4, value=mies).number_format = PLN
        ws.cell(row=r, column=5, value=naras).number_format = PLN
        ws.cell(row=r, column=6, value=f'=E{r}/C{r}').number_format = "0.0%"
    # BLAD: kwota jako tekst (wiersz 7 = dzial 750)
    ws["D7"] = f"{WYKONANIE['750'][7]:,.2f}".replace(",", " ").replace(".", ",") + " zł"
    # BLAD: scalona komorka w nazwie dzialu (wiersze 10-11 = 851 i 852)
    ws["B11"] = None
    ws.merge_cells("B10:B11")
    # BLAD: brak planu -> #DIV/0! w F13 (dzial 921 po przesunieciu = wiersz 13)
    ws["C13"] = None
    # BLAD: ukryta kolumna E
    ws.column_dimensions["E"].hidden = True
    # BLAD: niespojne daty
    ws["A17"], ws["A18"] = "Sporządzono: 05.09.2026", "Zatwierdzono: 2026-09-08"
    ws["A15"], ws["A15"].font = "Razem", BOLD
    for col in "CDE":
        ws[f"{col}15"] = f"=SUM({col}6:{col}14)"
        ws[f"{col}15"].number_format = PLN
    szerokosci(ws, [8, 42, 16, 20, 20, 16])
    wb.save(os.path.join(HERE, "zestawienie_bledy.xlsx"))


# --- sandbox do nawigacji (zadanie 2) ---------------------------------------
def sandbox():
    folder = os.path.join(HERE, "01_sandbox")
    if os.path.isdir(folder):
        shutil.rmtree(folder)
    os.makedirs(folder)

    def maly_xlsx(nazwa, miesiac_txt):
        wb = Workbook()
        ws = wb.active
        ws.title = "Zestawienie"
        ws.append(["Miesiąc", miesiac_txt])
        ws.append([])
        ws.append(["Dział", "Wykonanie"])
        for kod, _, plan in DZIALY[:3]:
            ws.append([kod, round(plan / 12 / 1000) * 1000])
        wb.save(os.path.join(folder, nazwa))

    for rok, mies in [(2025, 10), (2025, 11), (2025, 12)] + [(2026, m) for m in range(1, 7)]:
        maly_xlsx(f"zestawienie_{rok}_{mies:02d}.xlsx", f"{mies:02d}/{rok}")
    maly_xlsx("Kopia zestawienie_2026_03.xlsx", "03/2026")
    maly_xlsx("zestawienie_marzec_2026 (1).xlsx", "03/2026")
    maly_xlsx("zestawienie_2026_04 - poprawione.xlsx", "04/2026")
    for nazwa, tresc in [
        ("notatka_dla_banku_v2.txt", "Notatka dla banku – wersja 2 (robocza).\nWykonanie za I półrocze zgodne z planem.\n"),
        ("notatka_dla_banku_v3_FINAL.txt", "Notatka dla banku – wersja 3.\nWykonanie za I półrocze: 48,7% planu.\n"),
        ("notatka_dla_banku_v3_FINAL_poprawiona.txt", "Notatka dla banku – wersja 3 poprawiona.\nWykonanie za I półrocze: 48,9% planu.\n"),
        ("TODO.txt", "- wysłać zestawienie za czerwiec\n- poprawić prognozę dla banku\n- archiwum 2025\n"),
    ]:
        with open(os.path.join(folder, nazwa), "w", encoding="utf-8") as f:
            f.write(tresc)
    with open(os.path.join(folder, "eksport_erp_2026-06.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Data księgowania", "Dział", "Kwota"])
        for kod, _, plan in DZIALY:
            w.writerow(["2026-06-15", kod, f"{round(plan / 12):.2f}"])


if __name__ == "__main__":
    zapisz_erp("2026-09", brudny=True)
    zapisz_erp("2026-10", brudny=False)
    zestawienie_miesieczne_szablon()
    zestawienie_roczne()
    zestawienie_bledy()
    sandbox()
    print("Klucz odpowiedzi – sumy wg działów (bez duplikatu):")
    for m in ("2026-09", "2026-10"):
        print(f"  {m}: " + ", ".join(f"{k}={sum(p[4] for p in v):,}" for k, v in ERP[m].items())
              + f" | razem={sum(p[4] for v in ERP[m].values() for p in v):,}")
    print("Wykonanie I-VIII (roczne):")
    for kod, _, _ in DZIALY:
        print(f"  {kod}: {WYKONANIE[kod]} suma={sum(WYKONANIE[kod]):,}")
    print("Gotowe.")
